"""
Turn an uploaded file into text sections for indexing.

A section is one block of text plus an optional page number. PDFs give one section per
page, so chunks (and answer citations) can point back to the page they came from.
"""
from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass
from pathlib import Path

from ..models import Document

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 300
# Extracted text limit: about 75k tokens, a few cents of embeddings at most.
MAX_EXTRACTED_CHARS = 300_000

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"  # DOCX files are zip archives

EXTENSIONS = {
    ".pdf": Document.Source.PDF,
    ".docx": Document.Source.DOCX,
    ".txt": Document.Source.PLAIN,
    ".md": Document.Source.PLAIN,
    ".markdown": Document.Source.PLAIN,
}


class ExtractionError(ValueError):
    """The file cannot be indexed; the message is safe to show to the user."""


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None


def extract_sections(filename: str, data: bytes) -> tuple[str, list[Section]]:
    """Return (source type, non-empty sections) for an uploaded file."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError(f"File is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")

    source = EXTENSIONS.get(Path(filename).suffix.lower())
    if source is None:
        raise ExtractionError("Unsupported file type. Upload a PDF, DOCX, TXT or Markdown file.")

    # Check the content, not just the extension, before handing bytes to a parser.
    if source == Document.Source.PDF and not data.startswith(PDF_MAGIC):
        raise ExtractionError("This file is not a valid PDF.")
    if source == Document.Source.DOCX and not data.startswith(ZIP_MAGIC):
        raise ExtractionError("This file is not a valid DOCX document.")

    if source == Document.Source.PDF:
        sections = _pdf_sections(data)
    elif source == Document.Source.DOCX:
        sections = [Section(_docx_text(data))]
    else:
        sections = [Section(_plain_text(data))]

    sections = [s for s in sections if s.text.strip()]
    if not sections:
        hint = " Scanned PDFs need OCR, which is not supported." if source == Document.Source.PDF else ""
        raise ExtractionError("No text could be extracted from this file." + hint)
    if sum(len(s.text) for s in sections) > MAX_EXTRACTED_CHARS:
        raise ExtractionError(
            f"The file has more than {MAX_EXTRACTED_CHARS:,} characters of text. Split it into smaller files."
        )
    return source, sections


def _pdf_sections(data: bytes) -> list[Section]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("Password-protected PDFs are not supported.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractionError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")
        return [
            Section(text=(page.extract_text() or "").strip(), page=number)
            for number, page in enumerate(reader.pages, start=1)
        ]
    except ExtractionError:
        raise
    except (PyPdfError, ValueError, KeyError, TypeError) as e:
        logger.warning("PDF extraction failed: %s", e)
        raise ExtractionError("This PDF could not be read. It may be damaged.") from e


def _docx_text(data: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as e:
        logger.warning("DOCX extraction failed: %s", e)
        raise ExtractionError("This DOCX document could not be read. It may be damaged.") from e

    # Blank lines between blocks let the chunker keep paragraphs together.
    blocks = [p.text.strip() for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(b for b in blocks if b)


def _plain_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise ExtractionError("Text files must be UTF-8 encoded.") from e
