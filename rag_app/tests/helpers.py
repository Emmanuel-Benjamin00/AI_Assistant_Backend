"""Shared test fixtures: fake embeddings, tiny PDF/DOCX files and indexed chunks."""
from __future__ import annotations

import io

from ..models import Chunk, Document

DIMS = 1536


def fake_vector(seed: int) -> list[float]:
    v = [0.0] * DIMS
    v[seed % DIMS] = 1.0
    return v


def fake_embed(texts):
    return [fake_vector(len(t)) for t in texts]


def make_chunk(title: str, text: str, seed: int, **metadata) -> Chunk:
    doc, _ = Document.objects.get_or_create(title=title)
    return Chunk.objects.create(document=doc, text=text, embedding=fake_vector(seed), metadata=metadata)


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page (no PDF library needed)."""
    objects: list[bytes] = []
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    objects.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(len(pages)).encode() + b" >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("latin-1")
        stream = b"BT /F1 12 Tf 72 720 Td (" + escaped + b") Tj ET"
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            b"/Contents " + str(page_ids[i] + 1).encode() + b" 0 R >>"
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
